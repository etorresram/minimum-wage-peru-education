"""Construye notebook con código R y Word comentado desde los resultados en R.
Python solo prepara los documentos; las estimaciones y su revisión se hacen en R.
"""
from pathlib import Path
import base64, csv, json, subprocess, tempfile, uuid, math, re
from copy import deepcopy
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

OUT = Path(__file__).resolve().parent
SRC = (OUT/'feedback_coautor.R').read_text()
parts = {}
parts['setup'], rest = SRC.split('outcomes <- list(', 1)
parts['models'], rest = ('outcomes <- list('+rest).split('# Diagnóstico descriptivo', 1)
parts['support'], rest = ('# Diagnóstico descriptivo'+rest).split('set.seed(42)', 1)
parts['plot'] = ('set.seed(42)'+rest).split("cat('\\nRESULTADOS", 1)[0] + '\nprint(p)\n'

INTRO = ('Esta nota responde a la propuesta de concentrar el paper en salarios, formalidad y autoempleo, '
         'y de explorar una regresión discontinua por educación. Reestimamos los tres resultados con el diseño '
         'original y examinamos dos contrastes cerca del límite entre secundaria y educación superior. '
         'El DiD amplio muestra una caída relativa de la formalidad y un aumento del autoempleo, pero '
         'estas diferencias se atenúan en la comparación educativa más cercana. La información disponible '
         'no fundamenta una interpretación causal por regresión discontinua en ese límite.')
SAMPLE = ('Usamos cortes transversales repetidos de ENAHO entre 2021T1 y 2024T4, personas de 14 a 65 años '
          'con categoría educativa válida y ponderador positivo. Excluimos 2022T2 porque mezcla observaciones '
          'anteriores y posteriores al aumento de mayo. El período previo es 2021T1–2022T1 y el posterior '
          '2022T3–2024T4. No seguimos a las mismas personas en el tiempo.')
INFERENCE = ('Todas las regresiones usan mínimos cuadrados ponderados por fac500a, efectos fijos de departamento '
             'y trimestre calendario, y errores estándar agrupados por departamento. Los efectos de trimestre absorben '
             'los cambios nacionales comunes; los de departamento, diferencias territoriales persistentes. La agrupación '
             'permite correlación de los errores dentro de cada departamento a lo largo del tiempo. Hay 25 grupos en '
             'las nueve estimaciones. Los intervalos y valores p usan la corrección de muestra finita de fixest y una '
             'distribución t con 24 grados de libertad. Esta inferencia no corrige sesgos de identificación.')
VARS = [
 ('Salario real por hora', 'Logaritmo del salario horario en la ocupación principal entre asalariados privados con salario válido; incluye trabajadores del hogar. Se conserva la conversión de frecuencia de pago, el IPC del mes de referencia y la winsorización de la base original. No mide el ingreso de quienes salen del empleo asalariado.'),
 ('Formalidad', 'Indicador de empleo formal entre ocupados no públicos. El coeficiente multiplicado por 100 expresa puntos porcentuales de esa proporción; no es la probabilidad de tener empleo formal en toda la población en edad de trabajar.'),
 ('Autoempleo', 'Indicador de trabajador por cuenta propia, categoría p507 = 2, entre ocupados no públicos. Excluye empleadores. Formalidad y autoempleo no son categorías excluyentes: un trabajador por cuenta propia puede ser formal.')]
MAIN = ('Sea L igual a uno para educación baja, hasta secundaria completa (p301a de 1 a 6), y cero para '
        'alguna educación superior (7 a 11). P indica el período posterior. Para cada uno de los tres resultados '
        'estimamos la ecuación (1). Z contiene edad, edad al cuadrado, sexo, estado civil y residencia urbana; E '
        'son los años imputados a la categoría educativa. Los efectos fijos se representan por α del departamento '
        'y λ del trimestre. No se agrega P por separado porque ya está absorbido por λ.')
MAIN_INTERP = ('El coeficiente τ es el cambio posterior menos previo del grupo con educación baja, relativo al cambio '
               'del grupo con educación alta, ajustado por controles. La razón para mantener esta especificación es '
               'concentrar la discusión en los tres resultados pedidos sin cambiar simultáneamente el diseño del paper. '
               'Su interpretación causal exigiría tendencias paralelas condicionales en ausencia de la reforma, '
               'ausencia de shocks diferenciales coincidentes y comparabilidad de las muestras. La educación alta '
               'es un grupo de comparación de menor exposición promedio, no un grupo con exposición garantizada igual a cero.')
LOCAL = ('La ecuación (2) restringe la muestra a secundaria completa (p301a = 6, 11 años imputados) y superior '
         'no universitaria incompleta (p301a = 7, 13 años). Conserva los controles Z, pesos y efectos fijos. No '
         'incluye E por separado porque, con solo esas dos categorías, E = 13 − 2L y es perfectamente colineal '
         'con la constante y L. La restricción pregunta si el patrón del DiD amplio también aparece entre '
         'grupos educativos más cercanos. τ local sigue siendo una diferencia en diferencias; requiere sus propias '
         'tendencias paralelas y no elimina la selección hacia la educación superior.')
RD = ('Para explorar la sugerencia de RD fijamos c = 12 únicamente como punto intermedio entre las categorías '
      'de 11 y 13 años. Definimos x = E − 12 y A = 1(E ≥ 13). En la ecuación (3) usamos |x| ≤ 3, es decir, '
      'los valores imputados 9, 11, 13 y 14. Es la ventana simétrica entera más estrecha con al menos dos '
      'valores distintos por lado que permite estimar una pendiente a cada lado. No es un ancho de banda óptimo '
      'seleccionado por un estimador RD. Los 14 años agrupan superior no universitaria completa y universitaria incompleta.')
RD_INTERP = ('La ecuación permite cuatro pendientes: una a cada lado del límite en cada período. γ es el salto '
             'alto menos bajo previo; γ + θ es el salto posterior, evaluados en x = 0. Reportamos −θ para que todos '
             'los cuadros mantengan la orientación educación baja menos alta. El cálculo extrapola sobre el intervalo '
             'sin observaciones entre 11 y 13 años y depende de la forma lineal y de la codificación de E.')
RD_LIMIT = ('El límite c = 12 lo introduce el análisis; no documentamos una regla de asignación de la RMV que '
            'cambie al cruzarlo. Una RD requiere un mecanismo de asignación al tratamiento que cambie en un umbral '
            'conocido. Tener una variable discreta no invalida toda RD, pero aquí se añade un soporte muy escaso y '
            'años imputados a categorías. Por eso la ecuación (3) es una diferencia descriptiva de saltos, no una '
            'RD causal ni una RD difusa con un primer estadio estimado. Véase Cattaneo y Titiunik (2022).')
EXPOSURE = ('Como diagnóstico descriptivo, calculamos la fracción ponderada de asalariados privados bajo la RMV '
            'en el período previo. Es 37,63% para secundaria completa y 37,35% para superior no universitaria '
            'incompleta, una diferencia de 0,29 puntos porcentuales. Las medias son muy parecidas, pero no '
            'constituyen una prueba formal de ausencia de discontinuidad. Estar bajo la RMV tampoco equivale a '
            'recibir el tratamiento: intervienen horas trabajadas, cumplimiento y cobertura. Este gráfico no '
            'sirve para dividir el efecto de resultados por un supuesto primer estadio.')
CAUTIONS = [
 'El archivo de tendencias previas del paper rechaza la igualdad conjunta de los coeficientes anteriores a la reforma para salarios (p = 0,0013) y autoempleo (p = 0,0031); para formalidad p = 0,148. No rechazar en formalidad no demuestra tendencias paralelas, y estos diagnósticos del DiD amplio no validan automáticamente las muestras locales.',
 'La medición de formalidad utiliza el indicador oficial cuando está disponible y una reconstrucción en 2024. La comparación temporal hereda esa limitación. Los salarios y los dos indicadores de composición se condicionan en la ocupación observada, que también puede cambiar tras la reforma.',
 'Las caídas de formalidad y aumentos de autoempleo son cambios en proporciones de cortes transversales. No demuestran que las mismas personas abandonaron un empleo formal para convertirse en independientes. Para esa afirmación harían falta transiciones individuales identificadas.',
 'Como antecedente ya existente, la validación de la reforma de 2025 reporta autoempleo de signo opuesto (−1,43 puntos; p = 0,026). No añadimos una nueva batería de robustez: retenemos estas reservas porque afectan la interpretación de los tres resultados centrales.'
]
BRIEF = ('Al concentrar el análisis en salarios, formalidad y autoempleo, el DiD principal estima una variación '
         'relativa de −0,021 log puntos en el salario horario real (p = 0,106), una reducción de 4,64 puntos '
         'porcentuales en formalidad y un aumento de 3,63 puntos en autoempleo (ambos p < 0,001). En la comparación '
         'entre secundaria completa y superior no universitaria incompleta, las estimaciones son −0,025 log '
         'puntos, −1,40 y +1,30 puntos porcentuales, respectivamente, y sus intervalos al 95% incluyen cero. '
         'La comparación local no confirma con precisión el patrón amplio, pero tampoco prueba efectos nulos. '
         'La propuesta de discontinuidad queda como ejercicio exploratorio, porque el límite educativo utilizado '
         'no corresponde a una regla de asignación de la reforma. Las tendencias previas y la composición de '
         'las muestras requieren prudencia al atribuir los cambios a la RMV.')
EQ1 = r'Y_{irt}=\alpha_r+\lambda_t+\beta L_i+\tau(L_iP_t)+Z_i^{\prime}\Gamma+\kappa E_i+\varepsilon_{irt}'
EQ2 = r'Y_{irt}=\alpha_r+\lambda_t+\beta L_i+\tau_{local}(L_iP_t)+Z_i^{\prime}\Gamma+\varepsilon_{irt}'
EQ3 = r'\begin{aligned}Y_{irt}={}&\alpha_r+\lambda_t+\gamma A_i+b x_i+d(A_ix_i)\\&+q(P_tx_i)+h(A_iP_tx_i)+\theta(A_iP_t)+Z_i^{\prime}\Gamma+\varepsilon_{irt}\end{aligned}'

def md(text): return {'cell_type':'markdown','metadata':{},'id':uuid.uuid4().hex[:8],'source':text.splitlines(True)}
def code(text): return {'cell_type':'code','execution_count':None,'metadata':{},'id':uuid.uuid4().hex[:8],'source':text.splitlines(True),'outputs':[]}

cells = [md('# Salarios formalidad y autoempleo ante los comentarios del coautor\n\n'+INTRO+
 '\n\n**Cálculos en R.** Esta versión contiene celdas R nativas. Las celdas se ejecutaron secuencialmente con `Rscript` para guardar sus salidas; no se utilizó un kernel Jupyter en esta máquina. Para ejecución interactiva, seleccionar el kernel R (`IRkernel`). También se puede ejecutar `Rscript notebooks/feedback_coautor.R` desde `project/`. La base `data/processed/enaho_pooled.rds` debe estar disponible. Paquetes: `data.table`, `fixest`, `ggplot2`.'),
 md('## Muestra y variables\n\n'+SAMPLE+'\n\n'+'\n\n'.join('**'+a+'.** '+b for a,b in VARS)),
 code(parts['setup']+'\ncat("Observaciones de la muestra general:", nrow(d), "\\n")\nprint(d[, .(n=.N), by=.(post, skill)][order(post, skill)])\n'),
 md('## Especificación principal\n\n'+MAIN+'\n\n$$'+EQ1+r'\tag{1}$$'+'\n\n'+MAIN_INTERP+'\n\n'+INFERENCE),
 md('## Comparación educativa cercana\n\n'+LOCAL+'\n\n$$'+EQ2+r'\tag{2}$$'),
 md('## Exploración de la discontinuidad propuesta\n\n'+RD+'\n\n$$'+EQ3+r'\tag{3}$$'+'\n\n'+RD_INTERP+'\n\n'+RD_LIMIT+'\n\nReferencia: [Cattaneo y Titiunik 2022, Regression Discontinuity Designs](https://rdpackages.github.io/references/Cattaneo-Titiunik_2022_ARE.pdf).'),
 code(parts['models']),
 md('## Resultados y unidades\n\nUn coeficiente de −0,02065 en log salario equivale a una diferencia de aproximadamente −2,04% mediante 100(exp(τ)−1); no es una variación en soles. Los indicadores binarios se expresan en puntos porcentuales. Los coeficientes de los tres diseños están orientados como educación baja menos alta.\n\nLos CSV conservan los coeficientes originales en logaritmos o proporciones. Esta tabla convierte solo formalidad y autoempleo a puntos porcentuales.'),
 code('shown <- copy(res)\nshown[, unidad := ifelse(resultado == "Log salario real por hora", "log puntos", "pp")]\nfor (v in c("coef", "ee", "li95", "ls95"))\n  shown[unidad == "pp", (v) := get(v)*100]\nprint(shown[, .(diseno, resultado, n, unidad,\n                coef=round(coef,4), ee=round(ee,4),\n                li95=round(li95,4), ls95=round(ls95,4), p=signif(p,4))])\n'),
 md('La estimación amplia de salarios es imprecisa al 5%: su intervalo va aproximadamente de −4,50% a +0,47%. No equivale a demostrar ausencia de efecto. La formalidad y el autoempleo presentan asociaciones diferenciales en el DiD amplio. La comparación local reduce sus magnitudes y ofrece intervalos más amplios; el cambio del salto lineal tampoco se distingue de cero. Cambiar de muestra cambia la población comparada: la falta de significación local no es una prueba de que los coeficientes amplios y locales difieran estadísticamente.'),
 md('## Diagnóstico de exposición previo a la reforma\n\n'+EXPOSURE),
 code(parts['support']+'\nprint(support)\n'), code(parts['plot']),
 md('## Reservas que deben acompañar los tres resultados\n\n'+'\n\n'.join(CAUTIONS)),
 md('## Síntesis para que el coautor redacte\n\n'+BRIEF),
 md('## Revisión y trazabilidad\n\nLas nueve regresiones se contrastaron con `lm()` y una matriz de covarianza agrupada calculada a partir de los residuos y la matriz de diseño, sin reutilizar la matriz de covarianza de `fixest`. Los coeficientes y errores estándar difieren menos de 1e−9; los valores p menos de 1e−8. Las tres filas principales coinciden con `output/main_did_coefs.csv`. Se verificaron también el número de observaciones, los 25 departamentos, los signos y los intervalos.\n\nLa auditoría es `revisar_feedback_coautor.R`; sus resultados son `feedback_coautor_revision.csv`. `feedback_coautor_sesion.txt` registra versiones y el hash de la base. Los diagnósticos heredados provienen de `output/pretrend_tests.csv` y `output/validation2025_did.csv`. Para repetir la auditoría, ejecutar `Rscript notebooks/revisar_feedback_coautor.R` desde `project/`.\n\nLa réplica numérica comprueba los cálculos; no certifica los supuestos causales ni revalida toda la construcción de los microdatos.')]

# Ejecutar exactamente las celdas R en una sesión, capturando salidas por celda.
with tempfile.TemporaryDirectory(prefix='mw_notebook_') as tmp:
    tmp = Path(tmp)
    runner = ['options(width=120)\npdf(file=NULL)\n']
    cc = 0
    for i,c in enumerate(cells):
        if c['cell_type'] != 'code': continue
        cc += 1
        cellfile=tmp/f'cell_{i}.R'; cellfile.write_text(''.join(c['source']))
        outfile=tmp/f'output_{i}.txt'
        runner.append(f'sink({json.dumps(str(outfile))})\nsource({json.dumps(str(cellfile))}, local=.GlobalEnv, echo=FALSE)\nsink()\n')
        c['execution_count']=cc
    runner.append('dev.off()\n')
    (tmp/'run.R').write_text(''.join(runner))
    run=subprocess.run(['Rscript',str(tmp/'run.R')],cwd=OUT,text=True,capture_output=True,check=True)
    for i,c in enumerate(cells):
        if c['cell_type']!='code':continue
        text=(tmp/f'output_{i}.txt').read_text()
        if text: c['outputs'].append({'output_type':'stream','name':'stdout','text':text.splitlines(True)})
        if 'print(p)' in ''.join(c['source']):
            c['outputs'].append({'output_type':'display_data','metadata':{},'data':{'image/png':base64.b64encode((OUT/'feedback_coautor_exposicion.png').read_bytes()).decode(),'text/plain':'Exposición previa por categoría educativa'}})
    if run.stderr.strip(): print(run.stderr)
nb={'cells':cells,'metadata':{'kernelspec':{'display_name':'R','language':'R','name':'ir'},'language_info':{'name':'R','version':'4.6.1','file_extension':'.r','mimetype':'text/x-r-source'}},'nbformat':4,'nbformat_minor':5}
(OUT/'feedback_coautor.ipynb').write_text(json.dumps(nb,ensure_ascii=False,indent=1)+'\n')

# Word: contenido explicativo, ecuaciones nativas y comentarios anclados.
rows=list(csv.DictReader((OUT/'feedback_coautor_resultados.csv').open()))
doc=Document(); section=doc.sections[0]
section.page_width=Inches(8.5); section.page_height=Inches(11)
section.top_margin=section.bottom_margin=Inches(.7)
section.left_margin=section.right_margin=Inches(.8)
normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(11)
normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.08
for st in ['Title','Heading 1','Heading 2']:
 doc.styles[st].font.name='Calibri';doc.styles[st].font.color.rgb=RGBColor(0,0,0)
doc.styles['Title'].font.size=Pt(24)
doc.styles['Heading 1'].font.size=Pt(15)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(12)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(7)
doc.styles['Heading 2'].font.size=Pt(12)
doc.core_properties.title='Salarios formalidad y autoempleo ante los comentarios del coautor'
doc.core_properties.subject='Especificaciones econométricas y resultados de ENAHO 2021 a 2024'
doc.core_properties.author='Análisis para los coautores'
doc.core_properties.keywords='salario mínimo, ENAHO, DiD, educación'
comments=[]
def para(text,comment=None):
 p=doc.add_paragraph(text)
 if comment: comments.append((p,comment))
 return p

def head(text): return doc.add_heading(text,1)
def newpage(): doc.add_page_break()
def mathrun(text):
 r=OxmlElement('m:r');t=OxmlElement('m:t')
 t.set(qn('xml:space'),'preserve');t.text=text;r.append(t);return r
def sub(base,idx):
 s=OxmlElement('m:sSub');e=OxmlElement('m:e');e.append(mathrun(base));i=OxmlElement('m:sub');i.append(mathrun(idx));s.extend([e,i]);return s
def eq(elements,number=None):
 p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(9)
 om=OxmlElement('m:oMath')
 for x in elements: om.append(sub(*x) if isinstance(x,tuple) else mathrun(x))
 p._p.append(om)
 if number:p.add_run('   ('+str(number)+')')
 return p
base=[('Y','irt'),' = ',('α','r'),' + ',('λ','t')]

# Página 1
p=doc.add_paragraph('Salarios formalidad y autoempleo',style='Title')
para('Especificaciones y resultados para responder al coautor').runs[0].italic=True
para(INTRO, 'La propuesta de concentrar el artículo en estos tres resultados se implementa manteniendo una sola especificación principal. Los ejercicios locales responden a la sugerencia sobre educación, sin multiplicar pruebas de robustez.')
head('Datos y resultados de interés')
para(SAMPLE)
for title,text in VARS:
 p=doc.add_paragraph();p.add_run(title+'. ').bold=True;p.add_run(text)
 if title=='Formalidad':comments.append((p,'El denominador son los ocupados no públicos. Una caída de 4,64 puntos en esta proporción no significa una caída de 4,64 puntos en la probabilidad de empleo formal de toda la población.'))
para('Los años educativos son valores asignados a categorías de p301a. La base no contiene una medición continua de años aprobados. La estimación principal usa 69.534 observaciones para salarios y 181.317 para cada indicador de composición.')

# Página 2
newpage();head('Especificaciones econométricas')
para(MAIN)
eq(base+[' + β',('L','i'),' + τ(',('L','i'),('P','t'),') + ',('Z','i'),'′Γ + κ',('E','i'),' + ',('ε','irt')],1)
para(MAIN_INTERP)
para(LOCAL)
eq(base+[' + β',('L','i'),' + ',('τ','local'),'(',('L','i'),('P','t'),') + ',('Z','i'),'′Γ + ',('ε','irt')],2)
para(INFERENCE)

# Página 3
newpage();head('Qué permite evaluar el umbral educativo')
para(RD, 'El valor 12 es una normalización elegida entre dos categorías; no es un umbral institucional. No debe describirse como una regla que determina quién recibe el aumento de la RMV.')
eq(base+[' + γ',('A','i'),' + b',('x','i'),' + d(',('A','i'),('x','i'),')'])
eq([' + q(',('P','t'),('x','i'),') + h(',('A','i'),('P','t'),('x','i'),') + θ(',('A','i'),('P','t'),') + ',('Z','i'),'′Γ + ',('ε','irt')],3)
para(RD_INTERP)
para(RD_LIMIT, 'La discreción de una variable por sí sola no descarta una RD. El problema central aquí es la falta de una regla de asignación del tratamiento en el límite propuesto, sumada a la imputación de años y al escaso soporte.')
para(EXPOSURE, 'El porcentaje bajo la RMV es un diagnóstico descriptivo. No se estimó un primer estadio de una RD difusa ni se hizo una prueba formal de que la diferencia entre estas dos medias sea cero.')

# Página 4
newpage();head('Resultados de las nueve estimaciones')
para('Todos los signos expresan educación baja menos alta. El salario se reporta en log puntos; formalidad y autoempleo en puntos porcentuales. IC es el intervalo de confianza al 95%.')
table=doc.add_table(rows=1, cols=6);table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
widths=[1.04,1.0,.73,1.72,.63,.78]
for i,w in enumerate(widths):table.columns[i].width=Inches(w)
headers=['Resultado','Modelo','Coef.','IC 95%','p','N']
for cell,text in zip(table.rows[0].cells,headers):cell.text=text
for r in rows:
 cells_=table.add_row().cells
 binary=r['resultado']!='Log salario real por hora';scale=100 if binary else 1
 label={'Log salario real por hora':'Salario','Empleo formal':'Formalidad','Autoempleo':'Autoempleo'}[r['resultado']]
 design={'DiD principal: baja vs alta educación':'Principal','DiD local: 11 vs 13 años':'Local 11/13','Cambio del salto lineal en 12 años':'Salto lineal'}[r['diseno']]
 vals=[label,design,f"{float(r['coef'])*scale:+.3f}",f"[{float(r['li95'])*scale:.3f}, {float(r['ls95'])*scale:.3f}]",'<0,001' if float(r['p'])<.001 else f"{float(r['p']):.3f}",f"{int(r['n']):,}".replace(',','.')]
 for c,t in zip(cells_,vals):c.text=t.replace('.',',') if c not in (cells_[-1],) and t not in ('Principal','Local 11/13','Salto lineal','Salario','Formalidad','Autoempleo') else t
 # Restaurar N sin conversión decimal.
 cells_[-1].text=vals[-1]
for ri,row in enumerate(table.rows):
 for ci,cell in enumerate(row.cells):
  cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
  pr=cell._tc.get_or_add_tcPr();margins=OxmlElement('w:tcMar')
  for side in ['top','left','bottom','right']:
   e=OxmlElement('w:'+side);e.set(qn('w:w'),'70');e.set(qn('w:type'),'dxa');margins.append(e)
  pr.append(margins)
  shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DFE6EC' if ri==0 else ('F5F7F9' if ri%2==0 else 'FFFFFF'));pr.append(shade)
  borders=OxmlElement('w:tcBorders')
  for side in ['top','left','bottom','right']:
   el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
  pr.append(borders)
  for p in cell.paragraphs:
   p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3);p.paragraph_format.line_spacing=1
   p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci<2 else WD_ALIGN_PARAGRAPH.CENTER
   for run in p.runs:run.font.size=Pt(10);run.bold=(ri==0)
 trpr=table.rows[ri]._tr.get_or_add_trPr();el=OxmlElement('w:cantSplit');trpr.append(el)
repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
p=para('El salario del DiD amplio equivale a aproximadamente −2,04% con un intervalo de −4,50% a +0,47%. La estimación no permite afirmar que el efecto sea cero ni que sea un “nulo preciso”.', 'p = 0,106 implica falta de rechazo al 5%, no evidencia de equivalencia ni de ausencia de efecto. El intervalo es la guía para describir la precisión.')
p.paragraph_format.space_before=Pt(8)
para('La comparación local atenúa las diferencias de formalidad y autoempleo y produce intervalos que incluyen cero. El salto lineal también es impreciso. La pérdida de significación al restringir la muestra no demuestra que los coeficientes amplios y locales sean estadísticamente diferentes.')
head('Síntesis para la redacción')
para(BRIEF)

# Página 5
newpage();head('Exposición previa y alcance de la interpretación')
p=doc.add_paragraph();p.add_run().add_picture(str(OUT/'feedback_coautor_exposicion.png'),width=Inches(6.6))
p=para('Figura 1  Proporción de asalariados privados bajo la RMV antes de la reforma por categoría educativa. Los dos puntos de 14 años corresponden a categorías distintas; el desplazamiento horizontal mínimo solo facilita la lectura.');p.runs[0].font.size=Pt(9)
for i,text in enumerate(CAUTIONS[:3]):
 para(text, ('Estos diagnósticos provienen del análisis ya existente del paper. Se conservan por su relevancia para identificar el efecto, aunque se reduzca la cantidad de ejercicios de robustez.' if i==0 else 'El cambio de medición de formalidad en 2024 debe mencionarse al valorar la comparabilidad temporal.' if i==1 else 'La muestra es una sucesión de cortes transversales. Evitar escribir que los trabajadores “pasaron” de formalidad a autoempleo a partir de estos coeficientes.'))

# Página 6
newpage();head('Reproducibilidad y reservas finales')
para(CAUTIONS[3])
para('Todas las regresiones se calcularon en R 4.6.1 con fixest 0.14.2. Se revisaron con una segunda implementación mediante lm() y una matriz de covarianza agrupada construida a partir de la matriz de diseño y los residuos. En las nueve regresiones, las diferencias son inferiores a 1e−9 en coeficientes y errores estándar, y a 1e−8 en valores p. Las tres filas principales coinciden con los resultados guardados del paper. Esta revisión valida los cálculos, no los supuestos causales ni toda la preparación original de los microdatos.')
para('La notebook contiene las ecuaciones, código R visible y salidas calculadas. Para reproducirla interactivamente se necesita un kernel R de Jupyter; el script .R permite repetir las estimaciones desde RStudio o Rscript. La base procesada debe acompañar al entorno de réplica o generarse con el código del proyecto.')
files=[('feedback_coautor.ipynb','Notebook explicativa con código R y resultados.'),('feedback_coautor.R','Estimación de los nueve modelos y diagnóstico educativo.'),('feedback_coautor_resultados.csv','Coeficientes, errores estándar, valores p e intervalos.'),('revisar_feedback_coautor.R','Auditoría numérica independiente.'),('feedback_coautor_revision.csv','Diferencias entre ambas implementaciones.'),('feedback_coautor_sesion.txt','Versiones del entorno y hash de la base procesada.')]
for name,desc in files:
 p=doc.add_paragraph();p.add_run(name+'. ').bold=True;p.add_run(desc)
head('Fuentes')
para('ENAHO y archivos del proyecto: data/processed/enaho_pooled.rds; scripts/R/01_build_panel.R; scripts/R/03_main_estimation.R; output/main_did_coefs.csv; output/pretrend_tests.csv; output/validation2025_did.csv.')
para('Cattaneo, M. D. y Titiunik, R. (2022). Regression Discontinuity Designs. Annual Review of Economics, 14, 821–851. https://rdpackages.github.io/references/Cattaneo-Titiunik_2022_ARE.pdf')

# Comentarios reales de Word, anclados al texto; autor transparente.
for p,text in comments:
 doc.add_comment(p.runs,text=text,author='Revisión metodológica',initials='RM')
# Eliminar líneas decorativas heredadas del estilo Title de la plantilla.
for el in doc.styles.element.xpath('.//w:pBdr'):
 el.getparent().remove(el)
for el in doc.element.xpath('.//w:pBdr'):
 el.getparent().remove(el)
for shape in doc.inline_shapes:
 shape._inline.docPr.set('descr', 'Proporción previa de asalariados bajo la RMV según educación imputada. En 11 y 13 años las proporciones son 37,63% y 37,35%, respectivamente.')
# Números de página discretos en un documento largo.
footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
run=footer.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');run._r.addnext(fld)
path=OUT/'feedback_coautor_comentado.docx';doc.save(path)
print(f'Notebook: {len(cells)} celdas; {sum(c["cell_type"]=="code" for c in cells)} celdas R ejecutadas.')
print(f'Word: {len(comments)} comentarios anclados; {path}')
